#!/usr/bin/env python3
"""
gen_prebooking.py — อ่านไฟล์ PRE-WHEELCHAIR (แท็บรายวัน เช่น 01OCT26) → CSV ตาราง pas_wc_prebooking
(จองล่วงหน้าต่อไฟลท์: จำนวน WCHR/WCHS/WCHC/AVIH/MAAS แยกขาเข้า/ขาออก) ป้อน flow 'WC Build Cases' (wc-09)

โครงแท็บรายวัน:
  row0 col2 = วันที่ ("01 OCT 2026")
  row2 = หัว: 0 Airlines · 1 Flt no. · 2 Routing · 5 STA(time) · 6 STD(time) · 7 C/T OPEN · 8 C/T CLOSE
         Arrival: 9 WCHR · 10 WCHS · 11 WCHC · 12 AVIH · 13 MAAS · 14 TOTAL
         DEP:     15 WCHR · 16 WCHS · 17 WCHC · 18 AVIH · 19 MAAS

usage:  python3 tools/gen_prebooking.py  PRE-WHEELCHAIR.xlsx [อีกไฟล์...]  -o pas_wc_prebooking.csv
"""
import argparse, csv, datetime, re
from openpyxl import load_workbook

SKIP = {'DATA BASE', 'MASTER', 'Sheet1', '_CODES', 'ไม่ใช้'}
MON = {'JAN': 1, 'FEB': 2, 'MAR': 3, 'APR': 4, 'MAY': 5, 'JUN': 6,
       'JUL': 7, 'AUG': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DEC': 12}


def tab_date(ws, name):
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for c in row:
            if isinstance(c, datetime.datetime):
                return c.date()
            m = re.match(r'(\d{1,2})\s+([A-Z]{3})\s+(\d{4})', str(c or '').strip().upper())
            if m and m.group(2) in MON:
                return datetime.date(int(m.group(3)), MON[m.group(2)], int(m.group(1)))
    m = re.match(r'(\d{1,2})([A-Z]{3})(\d{2})', name.strip().upper())
    if m and m.group(2) in MON:
        return datetime.date(2000 + int(m.group(3)), MON[m.group(2)], int(m.group(1)))
    return None


def hhmm(v):
    if isinstance(v, datetime.time):
        return '%02d:%02d' % (v.hour, v.minute)
    if isinstance(v, datetime.datetime):
        return '%02d:%02d' % (v.hour, v.minute)
    m = re.match(r'^\s*(\d{1,2})[:.](\d{2})', str(v or ''))
    return '%02d:%02d' % (int(m.group(1)), int(m.group(2))) if m else ''


def n(v):
    try:
        return int(float(v))
    except Exception:
        return 0


def parse(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        name = ws.title.strip()
        if name in SKIP:
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((i for i, r in enumerate(rows[:6]) if r and len(r) > 1 and str(r[1]).strip() == 'Flt no.'), None)
        if hdr is None:
            continue
        d = tab_date(ws, name)
        if not d:
            continue
        for r in rows[hdr + 2:]:
            if not r or len(r) < 15:
                continue
            air = str(r[0] or '').strip()
            flt = str(r[1] or '').strip()
            if not air or not flt or air.upper() in ('TOTAL', 'AIRLINES'):
                continue
            if not re.match(r'^[A-Z0-9]{1,3}$', air):          # เฉพาะรหัสสายการบินจริง
                continue
            g = lambda i: n(r[i]) if len(r) > i else 0
            out.append({
                'Key': flt + '|' + d.isoformat(),
                'pas_bookingdate': d.isoformat(),
                'pas_flightno': flt, 'pas_airline': air, 'pas_routing': str(r[2] or '').strip(),
                'pas_sta': hhmm(r[5] if len(r) > 5 else ''), 'pas_std': hhmm(r[6] if len(r) > 6 else ''),
                'pas_ctopen': hhmm(r[7] if len(r) > 7 else ''), 'pas_ctclose': hhmm(r[8] if len(r) > 8 else ''),
                'arr_wchr': g(9), 'arr_wchs': g(10), 'arr_wchc': g(11), 'arr_avih': g(12), 'arr_maas': g(13),
                'dep_wchr': g(15), 'dep_wchs': g(16), 'dep_wchc': g(17), 'dep_avih': g(18), 'dep_maas': g(19),
            })
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('inputs', nargs='+', help='PRE-WHEELCHAIR (.xlsx) อย่างน้อย 1')
    ap.add_argument('-o', '--out', default='pas_wc_prebooking.csv')
    ap.add_argument('--nonzero', action='store_true', help='เก็บเฉพาะไฟลท์ที่มีจำนวน > 0 อย่างน้อย 1 บริการ')
    args = ap.parse_args()
    rows = []
    for p in args.inputs:
        r = parse(p)
        rows += r
        print(p, '→', len(r), 'rows')
    if args.nonzero:
        svc = ['arr_wchr', 'arr_wchs', 'arr_wchc', 'arr_avih', 'arr_maas', 'dep_wchr', 'dep_wchs', 'dep_wchc', 'dep_avih', 'dep_maas']
        rows = [x for x in rows if sum(x[k] for k in svc) > 0]
    if rows:
        with open(args.out, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
            w.writeheader(); w.writerows(rows)
    print('total', len(rows), '→', args.out)


if __name__ == '__main__':
    main()
