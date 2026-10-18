#!/usr/bin/env python3
"""
gen_porter_staff.py — นับ "เคสต่อพอตเตอร์ต่อวัน" จากไฟล์ PORTER SUMMARY (แท็บรายวัน เช่น 01OCT26)
→ CSV ตาราง pas_porter_staff (pas_name = Name_date · Work Date · Cases)

โครงแท็บรายวัน (DAILY PORTER SUMMARY):
  row0 col0 = วันที่ · หัวตารางมี 'ที่' ที่คอลัมน์ 1
  คอลัมน์ 0-based: 1 ที่(ลำดับ) · 2 สายการบิน · 3 เที่ยวบิน · 4 พนักงาน 1 · 5 พนักงาน 2 · 6 สถานะ
  แต่ละแถว = 1 งาน (เคส) · ชื่อในคอลัมน์ 4 และ/หรือ 5 = พอตเตอร์ที่รับเคสนั้น

usage:  python3 tools/gen_porter_staff.py  PORTER_SUMMARY.xlsx [อีกไฟล์...]  -o pas_porter_staff.csv
"""
import argparse, csv, datetime, re
from collections import defaultdict
from openpyxl import load_workbook

SKIP = {'MASTER ของแทร่', 'Sheet1', 'Sheet2', 'MONTHLY SUMMARY', 'ไม่ใช้', 'MASTER', '_CODES'}
MON = {'JAN': 1, 'FEB': 2, 'MAR': 3, 'APR': 4, 'MAY': 5, 'JUN': 6,
       'JUL': 7, 'AUG': 8, 'SEP': 9, 'OCT': 10, 'NOV': 11, 'DEC': 12}


def tab_date(ws, name):
    # 1) วันที่ใน row0 col0
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        for c in row[:2]:
            if isinstance(c, datetime.datetime):
                return c.date()
    # 2) จากชื่อแท็บ เช่น 01OCT26
    m = re.match(r'(\d{1,2})([A-Z]{3})(\d{2})', name.strip().upper())
    if m and m.group(2) in MON:
        return datetime.date(2000 + int(m.group(3)), MON[m.group(2)], int(m.group(1)))
    return None


def parse(path):
    wb = load_workbook(path, read_only=True, data_only=True)
    # cases[(name, date)] = จำนวนเคส
    cases = defaultdict(int)
    for ws in wb.worksheets:
        name = ws.title.strip()
        if name in SKIP:
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        hdr = next((i for i, r in enumerate(rows[:6]) if r and str(r[1]).strip() == 'ที่'), None)
        if hdr is None:
            continue
        d = tab_date(ws, name)
        if not d:
            continue
        for r in rows[hdr + 1:]:
            if not r or len(r) < 5:
                continue
            seq = str(r[1] or '').strip()
            if not re.match(r'^\d+$', seq):         # เฉพาะแถวงานจริง (มีลำดับที่)
                continue
            for col in (4, 5):                      # พนักงาน 1 / พนักงาน 2
                nm = str(r[col] or '').strip() if len(r) > col else ''
                if nm and nm not in ('-', 'None'):
                    cases[(nm, d)] += 1
    return cases


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('inputs', nargs='+', help='PORTER SUMMARY (.xlsx) อย่างน้อย 1')
    ap.add_argument('-o', '--out', default='pas_porter_staff.csv')
    args = ap.parse_args()
    total = defaultdict(int)
    for p in args.inputs:
        c = parse(p)
        for k, v in c.items():
            total[k] += v
        print(p, '→', len(c), 'porter-day rows')
    rows = sorted(total.items(), key=lambda kv: (kv[0][1], -kv[1]))
    with open(args.out, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f)
        w.writerow(['pas_name', 'pas_work_date', 'pas_emp_name', 'pas_cases'])
        for (nm, d), n in rows:
            w.writerow([nm + '_' + d.isoformat(), d.isoformat(), nm, n])
    print('total', len(rows), 'porter-day rows →', args.out)


if __name__ == '__main__':
    main()
