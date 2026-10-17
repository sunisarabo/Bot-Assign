#!/usr/bin/env python3
"""
extract_staffing.py — ดึงข้อมูลจากไฟล์วิเคราะห์การจัดพนักงาน (JAN-OCT) ออกเป็น CSV
  - flight_manning.csv   : คน/ไฟลท์จริงแยกเฟส CI/GATE/ARR ต่อทีม (จากชีต 'มาตรฐานไฟลท์')
  - shift_templates.csv  : เวลาเริ่มกะ/ความยาว/จำนวนคน ต่อทีม (จากชีต 'Template กะ')

usage:  python3 tools/extract_staffing.py  ANALYSIS.xlsx  [OUTDIR=tools]
ใช้ป้อน build_int_roster.py (flight_manning.csv) · shift_templates = แคตตาล็อกกะจริง
"""
import csv, os, re, statistics, sys
from collections import defaultdict
from openpyxl import load_workbook

# map ทีมในไฟล์วิเคราะห์ → ชื่อทีม ADV ในระบบ
TEAM_MAP = {'CHN': ['CHN'], 'EY': ['EY'], 'KE': ['KE'], 'PG': ['PG'], 'QR': ['QR'],
            'SQ': ['SQ'], 'TR': ['TR'], 'CHARTER': ['CHARTER'], 'PVT': ['PVT'],
            'JQ/TK': ['JQ', 'TK'], 'SU/AK': ['SU', 'AK'], 'WY/EK': ['EK'], 'WK/SV': ['WY/WK']}


def col(h, name):
    for i, c in enumerate(h):
        if c and name == str(c).strip():
            return i
    return -1


def flight_manning(wb, out):
    ws = wb['มาตรฐานไฟลท์']; rows = [list(r) for r in ws.iter_rows(values_only=True)]
    h = rows[3]
    iSup, iFC, iCnt, iGate, iArr, iDoc = (col(h, 'Supervisor'), col(h, 'Control/FC'), col(h, 'Counter'),
                                          col(h, 'Gate'), col(h, 'Arrival/Assist'), col(h, 'Doc/Crew'))
    num = lambda v: (float(v) if isinstance(v, (int, float)) else 0.0)
    agg = defaultdict(lambda: {'CI': [], 'GATE': [], 'ARR': []})
    for r in rows[4:]:
        if not r or not r[0]:
            continue
        t = str(r[0]).strip()
        agg[t]['CI'].append(num(r[iCnt]))
        agg[t]['GATE'].append(num(r[iSup]) + num(r[iFC]) + num(r[iGate]) + num(r[iDoc]))
        agg[t]['ARR'].append(num(r[iArr]))
    res = {}
    for ft, ph in agg.items():
        for mt in TEAM_MAP.get(ft, []):
            res[mt] = {k: round(statistics.mean(v), 2) if v else 0 for k, v in ph.items()}
    with open(os.path.join(out, 'flight_manning.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['Team', 'CI', 'GATE', 'ARR', 'Total'])
        for t in sorted(res):
            c, g, a = res[t]['CI'], res[t]['GATE'], res[t]['ARR']
            w.writerow([t, c, g, a, round(c + g + a, 1)])
    return len(res)


def shift_templates(wb, out):
    ws = wb['Template กะ']; rows = [list(r) for r in ws.iter_rows(values_only=True)]
    cur, outrows = None, []
    for r in rows:
        if not r or r[0] is None:
            continue
        c0 = str(r[0]).strip()
        m = re.match(r'^(.+?)\s*—\s*(ม\.ค\.|ต\.ค\.)', c0)
        if m:
            cur = (m.group(1).strip(), m.group(2)); continue
        if c0 == 'กะแนะนำ' and cur:
            try:
                st, ln, pp = int(r[1]), int(r[2]), int(r[3])
            except Exception:
                continue
            if pp > 0:
                outrows.append([cur[0], cur[1], st, ln, pp])
    with open(os.path.join(out, 'shift_templates.csv'), 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f); w.writerow(['Team', 'Period', 'StartHour', 'LenHrs', 'People'])
        w.writerows(outrows)
    return len(outrows)


def main():
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.dirname(os.path.abspath(__file__))
    wb = load_workbook(src, read_only=True, data_only=True)
    print('flight_manning.csv:', flight_manning(wb, out), 'teams')
    print('shift_templates.csv:', shift_templates(wb, out), 'shift rows')


if __name__ == '__main__':
    main()
